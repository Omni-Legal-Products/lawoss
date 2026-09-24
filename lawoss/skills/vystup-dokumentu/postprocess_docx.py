#!/usr/bin/env python3
"""Post-processing .docx: písmo kanceláře (výchozí Garamond 11 pt), okraje 2,5 cm, řádkování 1,15, tabulky s okraji.

Převzato z pracovního postupu RIHA legal (skill vystup-dokumentu, LAWOSS).
Písmo: proměnné prostředí LAWOSS_DOCX_FONT a LAWOSS_DOCX_SIZE.

Po úpravě .docx volitelně vyrenderuje věrné PDF přes LibreOffice (soffice).
PDF se generuje AŽ z hotového postprocessovaného .docx → PDF == .docx (vč. Garamondu).
Použití:
    postprocess_docx.py <soubor.docx>            # .docx + PDF (default)
    postprocess_docx.py <soubor.docx> --no-pdf   # jen .docx
"""
import sys
import os
import shutil
import subprocess
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH

FONT = os.environ.get("LAWOSS_DOCX_FONT", "Garamond")
SIZE = int(os.environ.get("LAWOSS_DOCX_SIZE", "11"))

def set_margins(doc):
    for s in doc.sections:
        s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5); s.right_margin = Cm(2.5)

def style_run(r):
    r.font.name = FONT
    r.font.size = Pt(SIZE)
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii','w:hAnsi','w:cs'):
        rfonts.set(qn(a), FONT)
    # zrušit kurzívu
    r.font.italic = False

def style_para(p, heading=False):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)
    for r in p.runs:
        style_run(r)
        if heading:
            r.font.bold = True
            r.font.size = Pt(13 if p.style.name.startswith('Title') else 12)

def left_align_hardbreak_paras(doc):
    """Past č. 3 (CLAUDE.md): odstavce s tvrdým zalomením (w:br bez w:type) zarovnat
    DOLEVA. V justifikované šabloně by jinak LibreOffice roztáhl každý řádek před
    zalomením přes celou šířku (adresy, identifikace stran, podpisy, label:value bloky)."""
    for p in doc.paragraphs:
        brs = list(p._p.iter(qn('w:br')))
        if any(b.get(qn('w:type')) is None for b in brs):
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

def fix_compact_numbered_paras(doc):
    """Past č. 8: pandoc dává jednopoložkovým („tight") číslovaným seznamům pStyle
    „Compact". LibreOffice u kombinace pStyle=Compact + přímé numPr číslování při
    renderu PDF zahodí (Word ne) — odstavec ztratí číslo. Oprava: pStyle odstranit,
    přímé numPr zůstává (loose položky pandocu pStyle také nemají)."""
    for p in doc.paragraphs:
        ppr = p._p.find(qn('w:pPr'))
        if ppr is None:
            continue
        style = ppr.find(qn('w:pStyle'))
        if style is not None and style.get(qn('w:val')) == 'Compact' \
                and ppr.find(qn('w:numPr')) is not None:
            ppr.remove(style)

def add_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4')
        e.set(qn('w:space'),'0'); e.set(qn('w:color'),'000000')
        borders.append(e)
    # tblBorders musí být ve schema pořadí PŘED shd/tblLayout/tblCellMar/tblLook,
    # jinak LibreOffice tblPr odmítne (Word je tolerantní)
    tblPr.insert_element_before(
        borders,
        'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
        'w:tblCaption', 'w:tblDescription',
    )

def find_soffice():
    """Najde binárku soffice na obvyklých cestách i v PATH."""
    for cand in (
        os.path.expanduser("~/.local/bin/soffice"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        shutil.which("soffice") or "",
        shutil.which("libreoffice") or "",
    ):
        if cand and os.path.exists(cand):
            return cand
    return None

def export_pdf(docx_path):
    """Vyrenderuje PDF z hotového .docx přes LibreOffice headless.

    Používá izolovaný profil (-env:UserInstallation), aby konverze nekolidovala
    s případně otevřenou GUI instancí LibreOffice a aby spolehlivě viděla
    font Garamond v ~/Library/Fonts. Vrací cestu k PDF nebo None.
    """
    soffice = find_soffice()
    if not soffice:
        print("VAROVÁNÍ: soffice nenalezen — PDF nevygenerováno (.docx je hotový).")
        return None
    outdir = os.path.dirname(os.path.abspath(docx_path)) or "."
    profile = "file://" + os.path.join(
        os.environ.get("TMPDIR", "/tmp").rstrip("/"), "lo_pdf_profile")
    cmd = [
        soffice, "--headless",
        f"-env:UserInstallation={profile}",
        "--convert-to", "pdf", "--outdir", outdir, docx_path,
    ]
    try:
        subprocess.run(cmd, check=True, stdout=subprocess.PIPE,
                       stderr=subprocess.PIPE, timeout=120)
    except subprocess.CalledProcessError as e:
        print(f"VAROVÁNÍ: konverze do PDF selhala: {e.stderr.decode(errors='replace')[:300]}")
        return None
    except subprocess.TimeoutExpired:
        print("VAROVÁNÍ: konverze do PDF vypršela (timeout 120 s).")
        return None
    pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
    if os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0:
        return pdf_path
    print("VAROVÁNÍ: PDF se nevytvořilo (prázdný/chybějící výstup).")
    return None

def main(path, make_pdf=True):
    doc = Document(path)
    set_margins(doc)
    for p in doc.paragraphs:
        is_h = p.style.name.startswith(('Heading','Title'))
        style_para(p, heading=is_h)
    left_align_hardbreak_paras(doc)
    fix_compact_numbered_paras(doc)
    for t in doc.tables:
        add_table_borders(t)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        style_run(r)
                        r.font.size = Pt(9)
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.0
    doc.save(path)
    print(f"OK: {path}")
    if make_pdf:
        pdf = export_pdf(path)
        if pdf:
            print(f"OK PDF: {pdf}")

if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if a != '--no-pdf']
    make_pdf = '--no-pdf' not in sys.argv
    if not args:
        sys.exit("Použití: postprocess_docx.py <soubor.docx> [--no-pdf]")
    main(args[0], make_pdf=make_pdf)
